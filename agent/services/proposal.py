"""Сборка коммерческого предложения из шаблона example.docx."""

from __future__ import annotations

import os
import re
import shutil
from copy import deepcopy
from datetime import datetime
from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table, _Row

from .catalog import Product, get_catalog

ROOT_DIR = Path(__file__).resolve().parents[2]
DEFAULT_TEMPLATE = Path(os.getenv("PROPOSAL_TEMPLATE_PATH") or ROOT_DIR / "assets" / "example.docx")
DEFAULT_OUTPUT_DIR = Path(os.getenv("OUTPUT_DIR") or ROOT_DIR / "output")

# Последний успешно сгенерированный файл (для отправки ботом)
LAST_PROPOSAL_PATH: str | None = None


def _format_money(value: float) -> str:
    return f"{value:,.2f}".replace(",", " ").replace(".", ",")


def _set_cell_text(cell, text: str) -> None:
    """Заменяет текст ячейки, сохраняя первый параграф по возможности."""
    text = "" if text is None else str(text)
    if cell.paragraphs:
        # очистить все runs первого параграфа
        p = cell.paragraphs[0]
        if p.runs:
            p.runs[0].text = text
            for run in p.runs[1:]:
                run.text = ""
        else:
            p.add_run(text)
        # удалить лишние параграфы
        for extra in cell.paragraphs[1:]:
            element = extra._element
            element.getparent().remove(element)
    else:
        cell.text = text


def _clone_row(table: Table, source_row: _Row) -> _Row:
    """Добавляет в конец таблицы копию строки (до вставки перед ИТОГО)."""
    new_tr = deepcopy(source_row._tr)
    table._tbl.append(new_tr)
    return table.rows[-1]


def _delete_row(table: Table, row: _Row) -> None:
    table._tbl.remove(row._tr)


def _find_items_table(doc: Document) -> Table:
    for table in doc.tables:
        if not table.rows:
            continue
        header = " ".join(c.text.strip().lower() for c in table.rows[0].cells)
        if "наименование" in header and ("№" in header or "п/п" in header or "пп" in header):
            return table
    raise ValueError("В шаблоне не найдена таблица позиций КП")


def _is_total_row(row: _Row) -> bool:
    joined = " ".join(c.text.strip().lower() for c in row.cells)
    return "итого" in joined


def _update_date_paragraph(doc: Document, when: datetime) -> None:
    date_str = when.strftime("%d.%m.%Y")
    pattern = re.compile(r"Ценовая информация\s+\d{2}\.\d{2}\.\d{4}", re.IGNORECASE)
    for paragraph in doc.paragraphs:
        if "Ценовая информация" in paragraph.text:
            new_text = pattern.sub(f"Ценовая информация {date_str}", paragraph.text)
            if new_text == paragraph.text and "Ценовая информация" in paragraph.text:
                new_text = f"Ценовая информация {date_str}"
            if paragraph.runs:
                paragraph.runs[0].text = new_text
                for run in paragraph.runs[1:]:
                    run.text = ""
            else:
                paragraph.text = new_text
            return


def build_proposal(
    items: list[dict[str, Any]],
    *,
    template_path: str | Path | None = None,
    output_dir: str | Path | None = None,
    when: datetime | None = None,
) -> dict[str, Any]:
    """
    items: [{"name": str, "qty": int}, ...]
    Возвращает dict с path, total, lines.
    """
    global LAST_PROPOSAL_PATH

    catalog = get_catalog()
    when = when or datetime.now()
    template_path = Path(template_path) if template_path else DEFAULT_TEMPLATE
    output_dir = Path(output_dir) if output_dir else DEFAULT_OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)

    if not items:
        raise ValueError("Список позиций пуст")

    resolved_lines: list[dict[str, Any]] = []
    ambiguities: list[dict[str, Any]] = []

    for raw in items:
        name = str(raw.get("name", "")).strip()
        qty = int(raw.get("qty", 1) or 1)
        if qty <= 0:
            raise ValueError(f"Некорректное количество для «{name}»: {qty}")
        product, candidates = catalog.resolve(name)
        if product is None:
            ambiguities.append(
                {
                    "requested": name,
                    "suggestions": [c.to_public_dict() for c in candidates],
                }
            )
            continue
        line_sum = product.price * qty
        resolved_lines.append(
            {
                "product": product,
                "qty": qty,
                "line_sum": line_sum,
            }
        )

    if ambiguities:
        return {
            "ok": False,
            "error": "Не удалось однозначно сопоставить позиции с прайсом",
            "ambiguities": ambiguities,
            "resolved": [
                {
                    "name": x["product"].name,
                    "qty": x["qty"],
                    "price": x["product"].price,
                }
                for x in resolved_lines
            ],
        }

    out_path = output_dir / f"KP_{when.strftime('%Y%m%d_%H%M%S')}.docx"
    shutil.copy(template_path, out_path)
    doc = Document(str(out_path))

    _update_date_paragraph(doc, when)
    table = _find_items_table(doc)

    # Разделяем строки: header / data / total
    header_row = table.rows[0]
    total_row = None
    data_rows: list[_Row] = []
    for idx, row in enumerate(table.rows):
        if idx == 0:
            continue
        if _is_total_row(row):
            total_row = row
        else:
            data_rows.append(row)

    if total_row is None:
        raise ValueError("В таблице КП не найдена строка ИТОГО")

    template_data_row = data_rows[0] if data_rows else None
    if template_data_row is None:
        raise ValueError("В таблице КП нет строки-образца для позиций")

    # Удаляем все data-строки, кроме первой (её используем как шаблон и перезапишем)
    for row in data_rows[1:]:
        _delete_row(table, row)

    # Нужно N строк перед ИТОГО. Сейчас осталась 1 data-строка.
    while len([r for r in table.rows[1:] if not _is_total_row(r)]) < len(resolved_lines):
        # вставляем копию перед total: append + move
        new_row = _clone_row(table, template_data_row)
        # переместить new_row перед total_row
        total_row._tr.addprevious(new_row._tr)

    # Если позиций меньше — удалить лишние data-строки (оставив ровно N)
    current_data = [r for r in table.rows[1:] if not _is_total_row(r)]
    while len(current_data) > len(resolved_lines):
        _delete_row(table, current_data[-1])
        current_data = [r for r in table.rows[1:] if not _is_total_row(r)]

    current_data = [r for r in table.rows[1:] if not _is_total_row(r)]
    total = 0.0
    public_lines: list[dict[str, Any]] = []

    for i, (row, line) in enumerate(zip(current_data, resolved_lines), start=1):
        product: Product = line["product"]
        qty = line["qty"]
        line_sum = line["line_sum"]
        total += line_sum
        cells = row.cells
        # № | Наименование | Описание | Кол-во | Цена | Сумма | Срок
        values = [
            str(i),
            product.name,
            product.description,
            str(qty),
            _format_money(product.price),
            _format_money(line_sum),
            "" if product.delivery_weeks is None else str(product.delivery_weeks),
        ]
        for cell, value in zip(cells, values):
            _set_cell_text(cell, value)
        public_lines.append(
            {
                "no": i,
                "name": product.name,
                "qty": qty,
                "price": product.price,
                "sum": line_sum,
                "delivery_weeks": product.delivery_weeks,
            }
        )

    # Обновить ИТОГО — сумма обычно в колонке «Сумма» (индекс 5)
    total_row = next(r for r in table.rows if _is_total_row(r))
    total_cells = total_row.cells
    if len(total_cells) >= 6:
        _set_cell_text(total_cells[5], _format_money(total))
    else:
        _set_cell_text(total_cells[-1], _format_money(total))

    # подавить unused warning
    _ = header_row

    doc.save(str(out_path))
    LAST_PROPOSAL_PATH = str(out_path.resolve())

    return {
        "ok": True,
        "path": LAST_PROPOSAL_PATH,
        "total": total,
        "total_formatted": _format_money(total),
        "lines": public_lines,
        "date": when.strftime("%d.%m.%Y"),
    }


def create_commercial_proposal(items_json: str) -> str:
    """Tool-friendly обёртка: items_json = '[{"name":"...","qty":1}]'."""
    import json

    try:
        items = json.loads(items_json)
    except json.JSONDecodeError as e:
        return json.dumps(
            {"ok": False, "error": f"Некорректный JSON: {e}"},
            ensure_ascii=False,
            indent=2,
        )

    if isinstance(items, dict) and "items" in items:
        items = items["items"]
    if not isinstance(items, list):
        return json.dumps(
            {"ok": False, "error": "Ожидается JSON-массив позиций"},
            ensure_ascii=False,
            indent=2,
        )

    try:
        result = build_proposal(items)
    except Exception as e:
        return json.dumps({"ok": False, "error": str(e)}, ensure_ascii=False, indent=2)

    return json.dumps(result, ensure_ascii=False, indent=2)


def pop_last_proposal_path() -> str | None:
    """Возвращает и сбрасывает путь последнего КП (для бота)."""
    global LAST_PROPOSAL_PATH
    path = LAST_PROPOSAL_PATH
    LAST_PROPOSAL_PATH = None
    return path
